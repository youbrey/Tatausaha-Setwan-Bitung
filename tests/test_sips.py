from __future__ import annotations

import hashlib
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from datetime import date
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from sekretariat_app.sips.models import InvitationFormData, TravelFormData
from sekretariat_app.sips.repository import SIPSRepository
from sekretariat_app.sips.service import SIPSService
from sekretariat_app.sips.sppd_generators import travel_cost_level


def assert_valid_docx(test_case: unittest.TestCase, path: Path) -> None:
    test_case.assertTrue(path.exists(), path)
    test_case.assertGreater(path.stat().st_size, 1_000, path)
    with zipfile.ZipFile(path) as archive:
        test_case.assertIsNone(archive.testzip(), path)
        test_case.assertIn("word/document.xml", archive.namelist())
        xml = b"".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
        test_case.assertNotIn(b"{{", xml, path)
        test_case.assertNotIn(b"{%", xml, path)


def all_docx_text(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue
            root = ET.fromstring(archive.read(name))
            chunks.extend(
                node.text or ""
                for node in root.iter()
                if node.tag.endswith("}t")
            )
    return "".join(chunks)


def table_grid_widths(table) -> list[int]:
    return [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst]


def row_height(row) -> int | None:
    properties = row._tr.trPr
    element = properties.find(qn("w:trHeight")) if properties is not None else None
    return int(element.get(qn("w:val"))) if element is not None else None


class SyntheticPersonnelMaster:
    """Data fiktif agar hasil test tidak memuat identitas personel kantor."""

    dprd = [
        {"nama": "ANGGOTA CONTOH A", "jabatan": "KETUA DPRD", "kategori": "Pimpinan DPRD"},
        {"nama": "ANGGOTA CONTOH B", "jabatan": "KETUA KOMISI I", "kategori": "Komisi I"},
        {"nama": "ANGGOTA CONTOH C", "jabatan": "ANGGOTA KOMISI I", "kategori": "Komisi I"},
        {"nama": "ANGGOTA CONTOH D", "jabatan": "ANGGOTA KOMISI II", "kategori": "Komisi II"},
    ]
    asn = [
        {
            "nama": "PEGAWAI CONTOH A", "nip": "000000000000000001",
            "pangkat": "PEMBINA", "jabatan": "SEKRETARIS DPRD",
        },
        {
            "nama": "PEGAWAI CONTOH B", "nip": "000000000000000002",
            "pangkat": "PENATA", "jabatan": "ANALIS CONTOH",
        },
    ]
    dprd_signers = ["KETUA DPRD - PENANDATANGAN CONTOH"]
    asn_signers = ["SEKRETARIS DPRD - PEGAWAI CONTOH A"]


class SIPSMigrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.master = SyntheticPersonnelMaster()
        cls.service = SIPSService(cls.master)

    def test_all_legacy_templates_are_packaged_byte_for_byte(self) -> None:
        project = Path(__file__).parents[1]
        legacy = project / "legacy" / "sips_app" / "resources" / "templates"
        packaged = project / "src" / "sekretariat_app" / "sips" / "resources" / "templates"
        legacy_files = {path.name: path for path in legacy.glob("*.docx")}
        packaged_files = {path.name: path for path in packaged.glob("*.docx")}
        self.assertEqual(len(legacy_files), 23)
        self.assertEqual(legacy_files.keys(), packaged_files.keys())
        for name, source in legacy_files.items():
            self.assertEqual(
                hashlib.sha256(source.read_bytes()).digest(),
                hashlib.sha256(packaged_files[name].read_bytes()).digest(),
                name,
            )

    def test_repository_can_list_only_matching_travel_drafts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            repository = SIPSRepository(Path(temp_dir) / "sips.sqlite3")
            common = dict(
                document_number="",
                document_date="2026-09-09",
                event_start="2026-09-10",
                event_end="2026-09-10",
                destination="Bitung",
                author="tester",
                payload={},
            )
            wanted = repository.save(
                record_type="travel_dprd", title="Draft DPRD",
                status="draft", **common,
            )
            repository.save(
                record_type="travel_dprd", title="Sudah dibuat",
                status="generated", **common,
            )
            repository.save(
                record_type="travel_secretariat", title="Draft Setwan",
                status="draft", **common,
            )

            drafts = repository.list(record_type="travel_dprd", status="draft")
            self.assertEqual([record.record_id for record in drafts], [wanted])

    def test_dprd_travel_generates_complete_document_set(self) -> None:
        data = TravelFormData(
            mode="dprd",
            document_numbers={
                "surat_tugas_dprd": "001/ST-DPRD/VIII/2026",
                "pemberitahuan_dprd": "002/PB-DPRD/VIII/2026",
                "spd_dprd": "003/SPD-DPRD/VIII/2026",
            },
            letter_date=date(2026, 8, 28),
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 2),
            travel_type="Kunjungan Kerja",
            basis_dprd="Keputusan Pimpinan DPRD Kota Bitung",
            basis_asn="Surat Perintah Sekretaris DPRD Kota Bitung",
            subject="Koordinasi penyusunan program kerja",
            notice_subject="Melaksanakan koordinasi penyusunan program kerja",
            destinations=["DPRD Kota Manado"],
            signer_dprd=self.master.dprd_signers[0],
            signer_asn=self.master.asn_signers[0],
            dprd=[self.master.dprd[0]],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_travel(data, temp_dir)
            self.assertEqual(len(files), 5)
            for path in files:
                assert_valid_docx(self, path)

    def test_travel_documents_match_reference_structure_for_multi_destination_trip(self) -> None:
        personnel = [
            {"nama": "ANGGOTA CONTOH A", "jabatan": "KETUA KOMISI I DPRD KOTA BITUNG", "kategori": "Komisi I"},
            {"nama": "ANGGOTA CONTOH B", "jabatan": "WAKIL KETUA KOMISI I DPRD KOTA BITUNG", "kategori": "Komisi I"},
            {"nama": "ANGGOTA CONTOH C", "jabatan": "SEKRETARIS KOMISI I DPRD KOTA BITUNG", "kategori": "Komisi I"},
            {"nama": "ANGGOTA CONTOH D", "jabatan": "ANGGOTA KOMISI III DPRD KOTA BITUNG", "kategori": "Komisi III"},
            {"nama": "ANGGOTA CONTOH E", "jabatan": "ANGGOTA KOMISI III DPRD KOTA BITUNG", "kategori": "Komisi III"},
        ]
        destinations = [
            "DPRD KABUPATEN MINAHASA UTARA",
            "DPRD KOTA MANADO",
            "DPRD KOTA TOMOHON",
        ]
        subject = (
            "Kunjungan Kerja ke DPRD Kabupaten Minahasa Utara, DPRD Kota Manado, "
            "dan DPRD Kota Tomohon dalam rangka Pengawasan DPRD terhadap "
            "Penggunaan Anggaran Perangkat Daerah"
        )
        data = TravelFormData(
            mode="dprd",
            document_numbers={
                "surat_tugas_dprd": "1235/SPT",
                "pemberitahuan_dprd": "12345/DPRD",
                "spd_dprd": "12345/SPD",
            },
            letter_date=date(2026, 9, 10),
            start_date=date(2026, 9, 10),
            end_date=date(2026, 9, 12),
            travel_type="Kunjungan Kerja",
            basis_dprd="Dasar surat contoh",
            basis_asn="",
            subject=subject,
            notice_subject=subject,
            destinations=destinations,
            signer_dprd=self.master.dprd_signers[0],
            signer_asn="",
            dprd=personnel,
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_travel(data, temp_dir)
            by_prefix = {
                prefix: next(path for path in files if path.name.startswith(prefix))
                for prefix in (
                    "surat-tugas-", "surat-pemberitahuan-", "spd-depan-", "daftar-hadir-",
                )
            }

            task = Document(by_prefix["surat-tugas-"])
            self.assertEqual(table_grid_widths(task.tables[0]), [704, 3544, 5380])
            self.assertEqual([row_height(row) for row in task.tables[0].rows], [597] + [510] * 5)

            notice = Document(by_prefix["surat-pemberitahuan-"])
            openings = [paragraph.text for paragraph in notice.paragraphs if "Bersama ini" in paragraph.text]
            self.assertEqual(len(openings), 3)
            display_destinations = [
                "DPRD Kabupaten Minahasa Utara", "DPRD Kota Manado", "DPRD Kota Tomohon",
            ]
            for index, (opening, destination) in enumerate(zip(openings, display_destinations)):
                self.assertIn(destination, opening)
                self.assertIn("dalam rangka Pengawasan DPRD", opening)
                self.assertNotIn("DPRD DPRD", opening)
                for other_index, other_destination in enumerate(display_destinations):
                    if index != other_index:
                        self.assertNotIn(other_destination, opening)

            attendance = Document(by_prefix["daftar-hadir-"])
            self.assertEqual(len(attendance.tables), 6)
            for table in attendance.tables[::2]:
                self.assertEqual(table_grid_widths(table), [704, 3969, 3119, 2126])
                self.assertEqual([row_height(row) for row in table.rows], [850] * 6)
                for row in table.rows[1:]:
                    self.assertEqual(row.cells[3].text, "")
                    self.assertEqual(len(row.cells[3].paragraphs), 1)
            attendance_text = "\n".join(paragraph.text for paragraph in attendance.paragraphs)
            self.assertEqual(attendance_text.count("DALAM RANGKA PENGAWASAN DPRD"), 6)
            self.assertNotIn("KE DPRD KABUPATEN MINAHASA UTARA, DPRD KOTA MANADO", attendance_text)

            spd = Document(by_prefix["spd-depan-"])
            self.assertEqual(len(spd.tables), 5)
            for table in spd.tables:
                self.assertEqual(table_grid_widths(table), [439, 3809, 3298, 2768])
                self.assertEqual(table.cell(8, 1).text.splitlines(), ["1.", "2.", "3."])
                self.assertEqual(table.cell(2, 2).paragraphs[2].text.strip(), "c. B")
                self.assertIn(
                    "Kunjungan Kerja ke DPRD Kabupaten Minahasa Utara, "
                    "DPRD Kota Manado, dan DPRD Kota Tomohon",
                    table.cell(3, 2).text,
                )
                destination_cell = " ".join(table.cell(5, 2).text.split())
                self.assertIn("Kabupaten Minahasa Utara, Kota Manado, Kota Tomohon", destination_cell)
                self.assertNotIn("dan Kota Tomohon", table.cell(5, 2).text)

    def test_travel_purpose_is_never_lost_when_form_fields_use_short_text(self) -> None:
        purpose = "Pengawasan DPRD terhadap Penggunaan Anggaran Perangkat Daerah"
        data = TravelFormData(
            mode="dprd",
            document_numbers={
                "surat_tugas_dprd": "010/ST-DPRD/IX/2026",
                "pemberitahuan_dprd": "020/PB-DPRD/IX/2026",
                "spd_dprd": "030/SPD-DPRD/IX/2026",
            },
            letter_date=date(2026, 9, 10),
            start_date=date(2026, 9, 10),
            end_date=date(2026, 9, 10),
            travel_type="Kunjungan Kerja",
            basis_dprd="Dasar surat pengujian",
            basis_asn="",
            subject=purpose,
            notice_subject=(
                "Pimpinan dan Anggota DPRD Kota Contoh akan melakukan "
                f"Kunjungan Kerja dalam rangka {purpose}"
            ),
            destinations=["DPRD KOTA MANADO"],
            signer_dprd=self.master.dprd_signers[0],
            signer_asn="",
            dprd=[self.master.dprd[1]],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_travel(data, temp_dir)
            notice_path = next(path for path in files if path.name.startswith("surat-pemberitahuan-"))
            attendance_path = next(path for path in files if path.name.startswith("daftar-hadir-"))
            spd_path = next(path for path in files if path.name.startswith("spd-depan-"))

            notice = Document(notice_path)
            notice_text = "\n".join(paragraph.text for paragraph in notice.paragraphs)
            self.assertIn(f"Kunjungan Kerja ke DPRD Kota Manado dalam rangka {purpose}", notice_text)

            attendance = Document(attendance_path)
            attendance_text = "\n".join(paragraph.text for paragraph in attendance.paragraphs)
            self.assertEqual(attendance_text.count(f"DALAM RANGKA {purpose.upper()}"), 2)

            spd = Document(spd_path)
            self.assertIn(f"dalam rangka {purpose}.", spd.tables[0].cell(3, 2).text)

    def test_travel_purpose_falls_back_to_other_form_field_after_dangling_marker(self) -> None:
        purpose = "Koordinasi Penguatan Tata Kelola Pemerintahan"
        data = TravelFormData(
            mode="dprd",
            document_numbers={
                "surat_tugas_dprd": "011/ST-DPRD/IX/2026",
                "pemberitahuan_dprd": "021/PB-DPRD/IX/2026",
                "spd_dprd": "031/SPD-DPRD/IX/2026",
            },
            letter_date=date(2026, 9, 10),
            start_date=date(2026, 9, 10),
            end_date=date(2026, 9, 10),
            travel_type="Kunjungan Kerja",
            basis_dprd="Dasar surat pengujian",
            basis_asn="",
            subject="Kunjungan Kerja ke DPRD Kota Manado dalam rangka",
            notice_subject=purpose,
            destinations=["DPRD Kota Manado"],
            signer_dprd=self.master.dprd_signers[0],
            signer_asn="",
            dprd=[self.master.dprd[1]],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_travel(data, temp_dir)
            task_path = next(path for path in files if path.name.startswith("surat-tugas-"))
            notice_path = next(path for path in files if path.name.startswith("surat-pemberitahuan-"))
            attendance_path = next(path for path in files if path.name.startswith("daftar-hadir-"))
            spd_path = next(path for path in files if path.name.startswith("spd-depan-"))

            task_text = "\n".join(paragraph.text for paragraph in Document(task_path).paragraphs)
            notice_text = "\n".join(paragraph.text for paragraph in Document(notice_path).paragraphs)
            attendance_text = "\n".join(paragraph.text for paragraph in Document(attendance_path).paragraphs)
            spd_text = Document(spd_path).tables[0].cell(3, 2).text
            for rendered in (task_text, notice_text, attendance_text, spd_text):
                self.assertIn(purpose.casefold(), rendered.casefold())

    def test_secretariat_travel_generates_executor_and_companion_sets(self) -> None:
        data = TravelFormData(
            mode="setwan",
            document_numbers={
                "surat_tugas_asn": "010/ST-SETWAN/VIII/2026",
                "pemberitahuan_asn": "011/PB-SETWAN/VIII/2026",
                "spd_pelaksana": "012/SPD-PL/VIII/2026",
                "spd_pendamping": "013/SPD-PD/VIII/2026",
            },
            letter_date=date(2026, 8, 28),
            start_date=date(2026, 9, 3),
            end_date=date(2026, 9, 4),
            travel_type="Kunjungan Konsultasi",
            basis_dprd="Keputusan Pimpinan DPRD Kota Bitung",
            basis_asn="Surat Perintah Sekretaris DPRD Kota Bitung",
            subject="Konsultasi tata kelola administrasi",
            notice_subject="Melaksanakan konsultasi tata kelola administrasi",
            destinations=["Sekretariat DPRD Kota Tomohon"],
            signer_dprd=self.master.dprd_signers[0],
            signer_asn=self.master.asn_signers[0],
            executors=[self.master.asn[0]],
            companions=[self.master.asn[1]],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_travel(data, temp_dir)
            self.assertEqual(len(files), 9)
            for path in files:
                assert_valid_docx(self, path)

    def test_dprd_mode_with_asn_only_does_not_create_dprd_attendance(self) -> None:
        data = TravelFormData(
            mode="dprd",
            document_numbers={
                "surat_tugas_asn": "020/ST-ASN/VIII/2026",
                "izin_pendamping": "020/IZIN-ASN/VIII/2026",
                "pemberitahuan_dprd": "021/PB-DPRD/VIII/2026",
                "spd_asn": "022/SPD-ASN/VIII/2026",
            },
            letter_date=date(2026, 8, 28),
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 1),
            travel_type="Kunjungan Kerja",
            basis_dprd="",
            basis_asn="Surat Perintah Sekretaris DPRD Kota Bitung",
            subject="Pendampingan konsultasi anggaran",
            notice_subject="Melaksanakan pendampingan konsultasi anggaran",
            destinations=["DPRD Kota Manado"],
            signer_dprd="",
            signer_asn=self.master.asn_signers[0],
            asn=[self.master.asn[0]],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_travel(data, temp_dir)
            self.assertEqual(len(files), 5)
            self.assertFalse(any("daftar-hadir" in path.name for path in files))

    def test_dprd_companion_documents_match_reference_wording_and_preview(self) -> None:
        destinations = [
            "DPRD Kabupaten Minahasa Utara",
            "DPRD Kota Manado",
            "DPRD Kota Tomohon",
        ]
        data = TravelFormData(
            mode="dprd",
            document_numbers={
                "surat_tugas_dprd": "100/ST-DPRD/IX/2026",
                "surat_tugas_asn": "101/ST-ASN/IX/2026",
                "izin_pendamping": "102/Setwan/IX/2026",
                "pemberitahuan_dprd": "103/PB-DPRD/IX/2026",
                "spd_dprd": "104/SPD-DPRD/IX/2026",
                "spd_asn": "105/SPD-ASN/IX/2026",
            },
            letter_date=date(2026, 9, 10),
            start_date=date(2026, 9, 10),
            end_date=date(2026, 9, 12),
            travel_type="Kunjungan Kerja",
            basis_dprd="Keputusan Pimpinan DPRD Kota Bitung",
            basis_asn="Surat Perintah Sekretaris DPRD Kota Bitung",
            subject="Pembahasan APBD Perubahan",
            notice_subject="Pembahasan APBD Perubahan",
            destinations=destinations,
            signer_dprd=self.master.dprd_signers[0],
            signer_asn=self.master.asn_signers[0],
            dprd=[self.master.dprd[1], self.master.dprd[2], self.master.dprd[3]],
            asn=[self.master.asn[1]],
        )
        expected_activity = (
            "Kunjungan Kerja Pimpinan dan Anggota Komisi I bersama Anggota "
            "Komisi II DPRD Kota Bitung ke DPRD Kabupaten Minahasa Utara, "
            "DPRD Kota Manado, dan DPRD Kota Tomohon dalam rangka "
            "Pembahasan APBD Perubahan"
        )

        preview = dict(self.service.travel_preview_documents(data))
        self.assertIn("permission-asn", preview)
        self.assertIn("surat-izin-pendamping", preview["permission-asn"])

        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_travel(data, temp_dir)
            task_path = next(path for path in files if path.name.startswith("surat-tugas-pendamping-"))
            permit_path = next(path for path in files if path.name.startswith("surat-izin-pendamping-"))
            spd_path = next(path for path in files if path.name.startswith("spd-pendamping-depan-"))
            for path in (task_path, permit_path, spd_path):
                assert_valid_docx(self, path)

            task = Document(task_path)
            task_text = all_docx_text(task_path)
            self.assertIn(f"Mendampingi {expected_activity}", task_text)
            self.assertNotIn("1. Nama", task_text)
            self.assertIn("Kepada", task_text)
            self.assertIn("PEMBINA", task_text)
            self.assertIn("NIP. 000000000000000001", task_text)
            self.assertEqual(len(task.tables), 1)

            spd = Document(spd_path)
            spd_name = spd.tables[0].cell(1, 2).text
            self.assertIn("PEGAWAI CONTOH B / 000000000000000002", spd_name)
            self.assertNotIn("\n", spd_name)
            self.assertEqual(spd.tables[0].cell(2, 2).paragraphs[2].text.strip(), "c. E")
            self.assertIn(f"Mendampingi {expected_activity}.", spd.tables[0].cell(3, 2).text)
            self.assertNotIn("Melakukan", spd.tables[0].cell(3, 2).text)

            permit_text = all_docx_text(permit_path)
            self.assertIn("102/Setwan/IX/2026", permit_text)
            self.assertIn(f"Dalam rangka Mendampingi {expected_activity}", permit_text)
            self.assertIn("ASISTEN PEREKONOMIAN DAN", permit_text)
            self.assertIn("PEGAWAI CONTOH B", permit_text)
            self.assertIn("NIP. 000000000000000001", permit_text)

            permit = Document(permit_path)
            closing = [
                paragraph.text for paragraph in permit.paragraphs
                if "Demikian permohonan" in paragraph.text
            ]
            self.assertEqual(
                closing,
                [
                    "Demikian permohonan ini kami sampaikan, atas perkenan "
                    "dan bantuannya diucapkan terima kasih."
                ],
            )

            notice_path = next(path for path in files if path.name.startswith("surat-pemberitahuan-"))
            notice = Document(notice_path)
            opening = next(paragraph.text for paragraph in notice.paragraphs if "Bersama ini" in paragraph.text)
            self.assertNotIn("Pendamping ASN", opening)
            self.assertNotIn("Staf Pendamping", opening)
            notice_text = "\n".join(paragraph.text for paragraph in notice.paragraphs)
            self.assertIn("Staf Pendamping", notice_text)

            attendance_path = next(path for path in files if path.name.startswith("daftar-hadir-"))
            attendance_text = "\n".join(
                paragraph.text for paragraph in Document(attendance_path).paragraphs
            )
            self.assertEqual(attendance_text.count("STAF PENDAMPING :"), len(destinations))
            self.assertEqual(attendance_text.count("PEGAWAI CONTOH B"), len(destinations))

    def test_spd_cost_level_follows_official_position_classification(self) -> None:
        self.assertEqual(travel_cost_level({}, dprd=True), "B")
        self.assertEqual(travel_cost_level({"jabatan": "Sekretaris DPRD"}), "C")
        self.assertEqual(
            travel_cost_level({"jabatan": "Kepala Bagian Umum dan Keuangan"}),
            "D",
        )
        for position in (
            "Kasubag Tata Usaha",
            "Staf Pelaksana",
            "Pengelola Peraturan Perundang - Undangan",
            "PPPK Teknis",
        ):
            self.assertEqual(travel_cost_level({"jabatan": position}), "E")

    def test_travel_validation_allows_optional_setwan_numbers(self) -> None:
        data = TravelFormData(
            mode="setwan",
            document_numbers={},
            letter_date=date(2026, 8, 28),
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 1),
            travel_type="Studi Komparasi",
            basis_dprd="",
            basis_asn="Surat Perintah Sekretaris DPRD Kota Bitung",
            subject="Studi komparasi administrasi persidangan",
            notice_subject="Melaksanakan studi komparasi administrasi persidangan",
            destinations=["Sekretariat DPRD Kota Tomohon"],
            signer_dprd="",
            signer_asn=self.master.asn_signers[0],
            executors=[self.master.asn[0]],
        )
        data.validate_preview()
        data.validate()

    def test_batch_report_keeps_other_files_when_one_generator_fails(self) -> None:
        data = TravelFormData(
            mode="dprd",
            document_numbers={
                "surat_tugas_dprd": "030/ST-DPRD/VIII/2026",
                "pemberitahuan_dprd": "031/PB-DPRD/VIII/2026",
                "spd_dprd": "032/SPD-DPRD/VIII/2026",
            },
            letter_date=date(2026, 8, 28),
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 1),
            travel_type="Kunjungan Kerja",
            basis_dprd="Keputusan Pimpinan DPRD Kota Bitung",
            basis_asn="",
            subject="Koordinasi penyusunan agenda",
            notice_subject="Melaksanakan koordinasi penyusunan agenda",
            destinations=["DPRD Kota Manado"],
            signer_dprd=self.master.dprd_signers[0],
            signer_asn="",
            dprd=[self.master.dprd[0]],
        )
        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "sekretariat_app.sips.service.buat_surat_tugas_dprd",
            side_effect=RuntimeError("template rusak"),
        ):
            report = self.service.generate_travel_report(data, temp_dir)
            self.assertEqual(len(report.failures), 1)
            self.assertIn("Surat Tugas DPRD", report.error_message)
            self.assertEqual(len(report.files), 4)
            for path in report.files:
                assert_valid_docx(self, path)

    def test_plenary_invitation_generates_eight_page_template_and_support(self) -> None:
        data = InvitationFormData(
            invitation_type="paripurna",
            number="080/UND-PAR/VIII/2026",
            letter_date=date(2026, 8, 28),
            meeting_date=date(2026, 9, 7),
            time_text="10.00 WITA s.d. selesai",
            agenda="Rapat Paripurna DPRD Kota Bitung",
            signer=self.master.dprd_signers[0],
            clothing="PSL",
            scenarios=["Pembukaan", "Penyampaian laporan", "Penutupan"],
            include_official_note=True,
            include_attendance=True,
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_invitation(data, temp_dir)
            self.assertEqual(len(files), 3)
            for path in files:
                assert_valid_docx(self, path)

    def test_regular_invitation_generates_supporting_documents(self) -> None:
        data = InvitationFormData(
            invitation_type="biasa",
            number="100/UND-DPRD/VIII/2026",
            letter_date=date(2026, 8, 28),
            meeting_date=date(2026, 9, 1),
            time_text="09.00 WITA s.d. selesai",
            agenda="Rapat pembahasan program kerja",
            signer=self.master.dprd_signers[0],
            meeting_executor="Pimpinan dan Anggota Komisi I",
            meeting_type="Rapat Kerja",
            related_parties=["Kepala Bagian Umum", "Tenaga Ahli Fraksi Nusantara"],
            other_destination_pages=[
                ["Wali Kota Bitung", "Sekretaris Daerah Kota Bitung"],
                ["Kepala Bappeda Kota Bitung"],
            ],
            include_official_note=True,
            include_attendance=True,
            include_related_attendance=True,
            include_secretariat_attendance=True,
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_invitation(data, temp_dir)
            self.assertEqual(len(files), 3)
            for path in files:
                assert_valid_docx(self, path)

    def test_supporting_documents_can_be_generated_independently(self) -> None:
        data = InvitationFormData(
            invitation_type="biasa",
            number="110/UND-DPRD/VIII/2026",
            letter_date=date(2026, 8, 28),
            meeting_date=date(2026, 9, 2),
            time_text="09.00 WITA s.d. selesai",
            agenda="Rapat evaluasi pelaksanaan kegiatan",
            signer=self.master.dprd_signers[0],
            meeting_executor="Pimpinan dan Anggota Komisi I",
            meeting_type="Rapat Kerja",
            related_parties=["Kepala Bagian Umum"],
            include_related_attendance=True,
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            note = self.service.generate_official_note(data, temp_dir)
            attendance = self.service.generate_meeting_attendance(data, temp_dir)
            assert_valid_docx(self, note)
            assert_valid_docx(self, attendance)

    def test_plenary_filename_uses_letter_date_like_legacy_sips(self) -> None:
        data = InvitationFormData(
            invitation_type="paripurna",
            number="005/DPRD/100/VIII/2026",
            letter_date=date(2026, 8, 31),
            meeting_date=date(2026, 9, 7),
            time_text="10.00 WITA s.d. selesai",
            agenda="Rapat Paripurna DPRD Kota Bitung",
            signer=self.master.dprd_signers[0],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            files = self.service.generate_invitation(data, temp_dir)
            self.assertEqual(files[0].name, "undangan-paripurna-senin-31-agustus.docx")

    def test_repository_supports_draft_recap_and_unique_numbers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            repository = SIPSRepository(Path(temp_dir) / "sekretariat.db")
            draft_id = repository.save(
                record_type="travel_dprd",
                title="Koordinasi",
                document_number="",
                document_date="2026-08-28",
                event_start="2026-09-01",
                event_end="2026-09-02",
                destination="Manado",
                status="draft",
                author="operator",
                payload={"mode": "dprd"},
            )
            generated_id = repository.save(
                record_type="invitation_regular",
                title="Rapat Kerja",
                document_number="100/UND/VIII/2026",
                document_date="2026-08-28",
                event_start="2026-09-01",
                event_end="2026-09-01",
                destination="Komisi I",
                status="generated",
                author="operator",
                payload={"invitation_type": "biasa"},
                numbers={"nomor_undangan": "100/UND/VIII/2026"},
            )

            self.assertEqual(repository.get(draft_id).status, "draft")
            self.assertEqual(repository.get(generated_id).status, "generated")
            self.assertEqual(len(repository.list(category="travel")), 1)
            self.assertEqual(len(repository.list(category="invitation", search="rapat")), 1)
            self.assertEqual(
                repository.dashboard_stats(),
                {"travel": 0, "plenary": 0, "regular": 1, "drafts": 1},
            )
            with self.assertRaisesRegex(ValueError, "sudah digunakan"):
                repository.validate_numbers({"nomor_undangan": "100/und/viii/2026"})
            repository.validate_numbers(
                {"nomor_undangan": "100/UND/VIII/2026"},
                generated_id,
            )
            with self.assertRaisesRegex(ValueError, "dua jenis dokumen"):
                repository.validate_numbers({"a": "NOMOR-1", "b": "nomor-1"})
            self.assertTrue(repository.delete_draft(draft_id))
            self.assertIsNone(repository.get(draft_id))

    def test_repository_detects_normalized_duplicate_travel_title(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            repository = SIPSRepository(Path(temp_dir) / "sekretariat.db")
            record_id = repository.save(
                record_type="travel_dprd",
                title="Koordinasi   Penyusunan APBD",
                document_number="001/ST/VIII/2026",
                document_date="2026-08-31",
                event_start="2026-09-01",
                event_end="2026-09-01",
                destination="Manado",
                status="generated",
                author="operator",
                payload={},
                numbers={"surat_tugas": "001/ST/VIII/2026"},
            )
            duplicate = repository.find_duplicate_travel_title(" koordinasi penyusunan apbd ")
            self.assertIsNotNone(duplicate)
            self.assertEqual(duplicate.record_id, record_id)
            self.assertIsNone(repository.find_duplicate_travel_title("Koordinasi Penyusunan APBD", record_id))

    def test_shell_uses_real_sips_pages_not_placeholders(self) -> None:
        shell_source = (
            Path(__file__).parents[1]
            / "src"
            / "sekretariat_app"
            / "ui"
            / "shell.py"
        ).read_text(encoding="utf-8")
        self.assertIn('TravelPage("dprd" if route == "travel_dprd" else "setwan"', shell_source)
        self.assertIn('InvitationPage("paripurna" if route == "invitation_plenary" else "biasa"', shell_source)
        self.assertNotIn("ModuleWorkspacePage(", shell_source)
        self.assertNotIn("= RecapPage(", shell_source)

    def test_travel_and_invitation_pages_use_automatic_live_preview(self) -> None:
        page_source = (
            Path(__file__).parents[1]
            / "src"
            / "sekretariat_app"
            / "ui"
            / "pages"
            / "sips.py"
        ).read_text(encoding="utf-8")
        preview_source = (
            Path(__file__).parents[1]
            / "src"
            / "sekretariat_app"
            / "ui"
            / "live_preview.py"
        ).read_text(encoding="utf-8")
        converter_source = (
            Path(__file__).parents[1]
            / "src"
            / "sekretariat_app"
            / "sips"
            / "preview.py"
        ).read_text(encoding="utf-8")

        self.assertEqual(page_source.count("self.live_preview = LiveDocumentPreview()"), 2)
        self.assertEqual(page_source.count("def _schedule_live_preview"), 2)
        self.assertNotIn('QPushButton("Pratinjau")', page_source)
        self.assertIn("class _LivePreviewWorker(QThread)", preview_source)
        self.assertIn("self._timer.setInterval(900)", preview_source)
        self.assertIn("$word.Visible=$false;$word.DisplayAlerts=0", converter_source)
        self.assertIn('"}finally{"', converter_source)
        self.assertIn('QPushButton("Buat Naskah Dinas Saja")', page_source)
        self.assertIn('QPushButton("Buat Daftar Hadir Saja")', page_source)
        self.assertIn("find_duplicate_travel_title", page_source)
        self.assertIn("QCompleter(DEFAULT_TRAVEL_DESTINATIONS", page_source)


if __name__ == "__main__":
    unittest.main()
