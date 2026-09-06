from __future__ import annotations

import html
import math
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz


class PasswordRequired(ValueError):
    def __init__(self, message: str, path: str):
        super().__init__(message)
        self.path = path


def open_pdf(path: str, password: str = "", *, owner: bool = False):
    document = fitz.open(path)
    if not document.is_pdf:
        document.close()
        raise ValueError("File yang dipilih bukan PDF.")
    encrypted = document.needs_pass or bool((document.metadata or {}).get("encryption"))
    if encrypted:
        access = document.authenticate(password)
        if not access or (owner and not access & 4):
            document.close()
            raise PasswordRequired(
                "Masukkan password pemilik PDF untuk mengedit atau melepas proteksi."
                if owner else "PDF memerlukan password yang benar.", path
            )
    return document


def page_numbers(value: str, count: int, *, repeat: bool = False) -> list[int]:
    """Parse halaman 1-based, rentang boleh terbalik untuk pengurutan."""
    if not value.strip():
        return list(range(count))
    result = []
    for part in value.split(","):
        match = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?", part)
        if not match:
            raise ValueError("Format halaman: 1,3,5-8. Nomor dimulai dari 1.")
        first, last = int(match[1]), int(match[2] or match[1])
        if min(first, last) < 1 or max(first, last) > count:
            raise ValueError(f"Nomor halaman harus antara 1 dan {count}.")
        step = 1 if last >= first else -1
        result.extend(range(first - 1, last - 1 + step, step))
        if len(result) > 100_000:
            raise ValueError("Daftar halaman terlalu panjang.")
    return result if repeat else list(dict.fromkeys(result))


def _save(document, target: str, **options) -> dict:
    if not len(document):
        raise ValueError("PDF harus mempunyai minimal satu halaman.")
    document.save(target, garbage=3, deflate=True, **options)
    return {"path": target, "count": len(document)}


def _visual_rect(page, values) -> fitz.Rect:
    if len(values) != 4 or not all(math.isfinite(float(v)) for v in values):
        raise ValueError("Area pilihan tidak valid.")
    x0, y0, x1, y1 = (max(0.0, min(1.0, float(v))) for v in values)
    rect = fitz.Rect(x0 * page.rect.width, y0 * page.rect.height,
                     x1 * page.rect.width, y1 * page.rect.height)
    if rect.width < 2 or rect.height < 2:
        raise ValueError("Tarik area yang lebih besar pada halaman.")
    return rect * page.derotation_matrix


def _blocks(page) -> list[dict]:
    blocks = []
    # Jangan mengekstrak byte gambar saat hanya mencari teks.
    raw = page.get_text("dict", flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)
    for block in raw["blocks"]:
        if block.get("type") != 0:
            continue
        spans = [s for line in block["lines"] for s in line["spans"]]
        text = "\n".join("".join(s["text"] for s in line["spans"]) for line in block["lines"])
        if not text.strip() or not spans:
            continue
        rect = fitz.Rect(block["bbox"]) * page.rotation_matrix
        first = spans[0]
        blocks.append({
            "id": block["number"], "text": text, "bbox": list(block["bbox"]),
            "visual": [rect.x0 / page.rect.width, rect.y0 / page.rect.height,
                       rect.x1 / page.rect.width, rect.y1 / page.rect.height],
            "size": first["size"], "font": first["font"],
            "color": f'#{first["color"]:06x}',
            "bold": bool(first["flags"] & 16), "italic": bool(first["flags"] & 2),
            "horizontal": all(abs(line.get("dir", (1, 0))[0] - 1) < .001
                              for line in block["lines"]),
        })
    return blocks


def render(request: dict) -> dict:
    with open_pdf(request["source"]) as document:
        index = min(max(0, int(request.get("page", 0))), len(document) - 1)
        page = document[index]
        scale = max(.1, min(3.0, float(request.get("width", 1000)) / page.rect.width))
        pixels = page.rect.width * page.rect.height * scale * scale
        if pixels > 2_000_000:
            scale *= math.sqrt(2_000_000 / pixels)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
        pixmap.save(request["output"])
        return {"path": request["output"], "page": index, "count": len(document),
                "width": page.rect.width, "height": page.rect.height, "blocks": _blocks(page)}


def edit(request: dict) -> dict:
    action = request["action"]
    with open_pdf(request["source"], request.get("password", ""), owner=True) as document:
        if action == "open":
            return _save(document, request["output"], encryption=fitz.PDF_ENCRYPT_NONE)
        if action == "organize":
            document.select(page_numbers(request["pages"], len(document), repeat=True))
        elif action == "delete":
            delete = set(page_numbers(request["pages"], len(document)))
            keep = [i for i in range(len(document)) if i not in delete]
            if not keep:
                raise ValueError("Tidak dapat menghapus seluruh halaman.")
            document.select(keep)
        elif action == "rotate":
            for index in page_numbers(request["pages"], len(document)):
                page = document[index]
                page.set_rotation((page.rotation + int(request["angle"])) % 360)
        elif action == "blank":
            index = int(request.get("page", 0))
            size = document[index].rect
            document.new_page(pno=index + 1, width=size.width, height=size.height)
        elif action in {"text", "replace_text", "image"}:
            page = document[int(request["page"])]
            rect = _visual_rect(page, request["rect"])
            if action == "image":
                from PIL import Image, ImageOps
                import io
                with Image.open(request["image"]) as original:
                    original.draft("RGB", (4096, 4096))
                    image = ImageOps.exif_transpose(original).convert("RGBA")
                    image.thumbnail((4096, 4096), Image.Resampling.LANCZOS)
                    data = io.BytesIO()
                    image.save(data, format="PNG")
                page.insert_image(rect, stream=data.getvalue(), keep_proportion=True,
                                  rotate=(-page.rotation) % 360)
            else:
                if action == "replace_text":
                    blocks = _blocks(page)
                    block = next((b for b in blocks if b["id"] == request["block_id"]), None)
                    if block is None or block["text"] != request.get("original"):
                        raise ValueError("Teks berubah. Pilih ulang blok pada preview terbaru.")
                    if not block["horizontal"]:
                        raise ValueError("Blok teks miring belum dapat diganti. Gunakan Tambah Teks.")
                    # Redaksi milik pengguna yang sudah ada tidak boleh ikut diterapkan.
                    if any(a.type[0] == fitz.PDF_ANNOT_REDACT for a in (page.annots() or [])):
                        raise ValueError("Halaman memiliki anotasi redaksi yang belum diterapkan.")
                    rect = fitz.Rect(block["bbox"])
                    if any(b["id"] != block["id"] and not (rect & fitz.Rect(b["bbox"])).is_empty for b in blocks):
                        raise ValueError("Blok bertumpuk dengan teks lain. Penggantian dibatalkan untuk menjaga teks di sekitarnya.")
                    page.add_redact_annot(rect, fill=False, cross_out=False)
                    page.apply_redactions(images=0, graphics=0, text=0)
                    if not request.get("text", "").strip():
                        return _save(document, request["output"], encryption=fitz.PDF_ENCRYPT_NONE)
                size = max(4, min(144, float(request.get("size", 12))))
                color = request.get("color", "#000000")
                if not re.fullmatch(r"#[0-9a-fA-F]{6}", color):
                    raise ValueError("Warna teks tidak valid.")
                family = {"serif": "serif", "sans-serif": "sans-serif", "monospace": "monospace"}.get(
                    request.get("family"), "sans-serif")
                css = (f"* {{font-family:{family};font-size:{size}pt;color:{color};}}"
                       f"body {{margin:0;line-height:1.1;font-weight:{'bold' if request.get('bold') else 'normal'};"
                       f"font-style:{'italic' if request.get('italic') else 'normal'};}}")
                text = html.escape(request.get("text", "")).replace("\n", "<br>")
                if not text.strip():
                    raise ValueError("Isi teks pengganti terlebih dahulu.")
                # Tidak diam-diam mengecilkan teks apabila kotak terlalu kecil.
                spare, _ = page.insert_htmlbox(rect, text, css=css, scale_low=1,
                    rotate=0 if action == "replace_text" else (-page.rotation) % 360)
                if spare < 0:
                    raise ValueError("Teks tidak muat. Kurangi ukuran huruf atau perbesar area. Perubahan dibatalkan.")
        elif action == "protect":
            user = request.get("user_password", "")
            owner = request.get("owner_password", "")
            if not user or not owner or user == owner:
                raise ValueError("Isi password buka dan password pemilik yang berbeda.")
            if max(len(user.encode()), len(owner.encode())) > 40:
                raise ValueError("Password maksimal 40 byte UTF-8.")
            permissions = fitz.PDF_PERM_ACCESSIBILITY
            if request.get("allow_print", True):
                permissions |= fitz.PDF_PERM_PRINT | fitz.PDF_PERM_PRINT_HQ
            if request.get("allow_copy", False):
                permissions |= fitz.PDF_PERM_COPY
            return _save(document, request["output"], encryption=fitz.PDF_ENCRYPT_AES_256,
                         owner_pw=owner, user_pw=user, permissions=permissions)
        elif action != "unprotect":
            raise ValueError(f"Operasi tidak dikenal: {action}")
        return _save(document, request["output"], encryption=fitz.PDF_ENCRYPT_NONE)


def merge(request: dict, progress) -> dict:
    with fitz.open() as result:
        sources = request["sources"]
        for index, source in enumerate(sources):
            with open_pdf(source, request.get("passwords", {}).get(source, ""), owner=True) as document:
                result.insert_pdf(document)
            progress(index + 1, len(sources), Path(source).name)
        return _save(result, request["output"])


def split(request: dict, progress) -> dict:
    directory = Path(request["output"])
    directory.mkdir(parents=True)
    with open_pdf(request["source"]) as document:
        groups = request.get("groups", "").strip()
        ranges = [page_numbers(g, len(document)) for g in groups.split(";")] if groups else [[i] for i in range(len(document))]
        for index, pages in enumerate(ranges):
            with fitz.open() as result:
                for page in pages:
                    result.insert_pdf(document, from_page=page, to_page=page)
                _save(result, str(directory / f"bagian-{index + 1:03d}.pdf"))
            progress(index + 1, len(ranges), "Memisahkan halaman")
    return {"path": str(directory), "files": len(ranges)}


def images_to_pdf(request: dict, progress) -> dict:
    from PIL import Image, ImageOps
    import io
    with fitz.open() as result:
        for index, path in enumerate(request["sources"]):
            with Image.open(path) as source:
                source.draft("RGB", (2560, 3508))
                oriented = ImageOps.exif_transpose(source)
                oriented.thumbnail((2560, 3508), Image.Resampling.LANCZOS)
                if oriented.mode in {"RGBA", "LA"} or "transparency" in oriented.info:
                    rgba = oriented.convert("RGBA")
                    photo = Image.new("RGB", rgba.size, "white")
                    photo.paste(rgba, mask=rgba.getchannel("A"))
                else:
                    photo = oriented.convert("RGB")
                buffer = io.BytesIO()
                photo.save(buffer, "JPEG", quality=95)
                page = result.new_page(width=595.276, height=841.89)
                page.insert_image(page.rect + (28.35, 28.35, -28.35, -28.35),
                                  stream=buffer.getvalue(), keep_proportion=True)
            progress(index + 1, len(request["sources"]), "Gambar ke PDF A4")
        return _save(result, request["output"])


def pdf_to_images(request: dict, progress) -> dict:
    directory = Path(request["output"])
    directory.mkdir(parents=True)
    with open_pdf(request["source"]) as document:
        pages = page_numbers(request.get("pages", ""), len(document))
        dpi = int(request.get("dpi", 150))
        for index, number in enumerate(pages):
            page = document[number]
            scale = dpi / 72
            if page.rect.width * page.rect.height * scale * scale > 40_000_000:
                raise ValueError("Ukuran gambar melebihi 40 megapiksel. Pilih DPI lebih kecil.")
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
            pixmap.save(str(directory / f"halaman-{number + 1:04d}.png"))
            del pixmap
            progress(index + 1, len(pages), "PDF ke PNG")
    return {"path": str(directory), "files": len(pages)}


def pdf_to_word(request: dict, progress) -> dict:
    """DOCX editable berbasis urutan blok, atau halaman sebagai gambar untuk fidelitas."""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    import io
    output = Document()
    image_mode = request.get("mode") == "appearance"
    text_pages = 0
    embedded_bytes = 0
    text_characters = 0
    with open_pdf(request["source"]) as document:
        for index, page in enumerate(document):
            section = output.sections[0] if index == 0 else output.add_section(WD_SECTION.NEW_PAGE)
            section.page_width, section.page_height = Pt(page.rect.width), Pt(page.rect.height)
            margin = 0 if image_mode else 20
            section.top_margin = section.bottom_margin = Pt(margin)
            section.left_margin = section.right_margin = Pt(margin)
            section.header_distance = section.footer_distance = Pt(0)
            if image_mode:
                scale = min(1.5, math.sqrt(2_000_000 / max(1, page.rect.width * page.rect.height)))
                pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False, colorspace=fitz.csRGB)
                image_bytes = pix.tobytes("jpeg", jpg_quality=95)
                embedded_bytes += len(image_bytes)
                if embedded_bytes > 96 * 1024 * 1024:
                    raise ValueError("Gambar DOCX melebihi anggaran memori 96 MB. Pisahkan PDF menjadi beberapa bagian sebelum konversi.")
                paragraph = output.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(1)
                run = paragraph.add_run()
                # Inline image menyisakan tinggi baris. Jangkar halaman tidak.
                inline = run.add_picture(io.BytesIO(image_bytes), width=Pt(page.rect.width), height=Pt(page.rect.height))._inline
                anchor = OxmlElement("wp:anchor")
                for key, value in {"distT":"0", "distB":"0", "distL":"0", "distR":"0", "simplePos":"0", "relativeHeight":"0", "behindDoc":"1", "locked":"0", "layoutInCell":"1", "allowOverlap":"1"}.items():
                    anchor.set(key, value)
                simple = OxmlElement("wp:simplePos"); simple.set("x", "0"); simple.set("y", "0"); anchor.append(simple)
                for axis in ("H", "V"):
                    position = OxmlElement(f"wp:position{axis}"); position.set("relativeFrom", "page")
                    offset = OxmlElement("wp:posOffset"); offset.text = "0"; position.append(offset); anchor.append(position)
                for tag in ("extent", "effectExtent"):
                    element = inline.find(qn(f"wp:{tag}"))
                    if element is not None: anchor.append(element)
                anchor.append(OxmlElement("wp:wrapNone"))
                for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
                    element = inline.find(qn(tag))
                    if element is not None: anchor.append(element)
                inline.getparent().replace(inline, anchor)
                del pix
            else:
                blocks = _blocks(page)
                text_pages += bool(blocks)
                if not blocks:
                    output.add_paragraph("[Halaman scan/gambar: tidak tersedia lapisan teks. Gunakan mode tampilan atau OCR eksternal.]")
                for block in blocks:
                    text_characters += len(block["text"])
                    if text_characters > 2_000_000:
                        raise ValueError("Teks terlalu besar untuk satu DOCX. Pisahkan PDF menjadi beberapa bagian sebelum konversi.")
                    paragraph = output.add_paragraph()
                    run = paragraph.add_run(block["text"])
                    run.font.size = Pt(block["size"])
                    run.font.bold, run.font.italic = block["bold"], block["italic"]
                    run.font.name = block["font"].split("+")[-1]
            progress(index + 1, len(document), "PDF ke Word")
    if not image_mode and not text_pages:
        raise ValueError("PDF tidak memiliki teks yang dapat diekstrak. Gunakan mode tampilan; OCR belum disertakan.")
    output.save(request["output"])
    return {"path": request["output"], "note": "Mode tampilan: halaman berupa gambar." if image_mode else "DOCX teks dapat diedit; tata letak kompleks dan tabel tidak direkonstruksi identik."}


def pdf_to_excel(request: dict, progress) -> dict:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    book = Workbook(write_only=True)
    found = 0
    with open_pdf(request["source"]) as document:
        for index, page in enumerate(document):
            finder = page.find_tables()
            for number, table in enumerate(finder.tables):
                sheet = book.create_sheet(f"H{index + 1}-T{number + 1}")
                for row in table.extract():
                    cells = []
                    for value in row:
                        cell = WriteOnlyCell(sheet, value=str(value or ""))
                        cell.data_type = "s"  # Teks PDF tidak dieksekusi sebagai formula Excel.
                        cells.append(cell)
                    sheet.append(cells)
                found += 1
            progress(index + 1, len(document), "Mendeteksi tabel")
    if not found:
        raise ValueError("Tidak ditemukan tabel teks. PDF scan atau tabel tanpa struktur memerlukan OCR/penyesuaian manual.")
    book.save(request["output"])
    return {"path": request["output"], "note": f"{found} tabel diekstrak. Periksa sel gabungan dan format angka."}


def office_to_pdf(request: dict, progress) -> dict:
    source = Path(request["source"]).resolve()
    target = Path(request["output"]).resolve()
    excel = source.suffix.lower() in {".xlsx", ".xls"}
    if os.name == "nt":
        # Argumen path melalui environment, bukan interpolasi shell.
        env = {**os.environ, "PDF_INPUT_PATH": str(source), "PDF_OUTPUT_PATH": str(target)}
        script = (
            "$ErrorActionPreference='Stop'; $app=$null; $doc=$null; try {"
            f"$app=New-Object -ComObject {'Excel' if excel else 'Word'}.Application;"
            "$app.Visible=$false; $app.DisplayAlerts=0; $app.AutomationSecurity=3;"
            + ("$app.AskToUpdateLinks=$false; $doc=$app.Workbooks.Open($env:PDF_INPUT_PATH,0,$true);"
               "$doc.ExportAsFixedFormat(0,$env:PDF_OUTPUT_PATH);" if excel else
               "$doc=$app.Documents.Open($env:PDF_INPUT_PATH,$false,$true);"
               "$doc.ExportAsFixedFormat($env:PDF_OUTPUT_PATH,17);") +
            "} finally {if($doc){$doc.Close(0);[void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($doc)}"
            "if($app){$app.Quit();[void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($app)}}"
        )
        try:
            subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
                           env=env, check=True, capture_output=True, timeout=180,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if target.is_file():
                return {"path": str(target)}
        except (OSError, subprocess.SubprocessError):
            target.unlink(missing_ok=True)
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        for root in (os.environ.get("PROGRAMFILES", ""), os.environ.get("PROGRAMFILES(X86)", "")):
            candidate = Path(root) / "LibreOffice/program/soffice.exe"
            if candidate.is_file():
                office = str(candidate); break
    if not office:
        raise RuntimeError("Konversi ini memerlukan Microsoft Word/Excel atau LibreOffice yang terpasang lokal.")
    with tempfile.TemporaryDirectory(prefix="office-pdf-") as tmp:
        profile = (Path(tmp) / "profile").as_uri()
        subprocess.run([office, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(source)],
                       check=True, capture_output=True, timeout=180,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        converted = Path(tmp) / f"{source.stem}.pdf"
        if not converted.is_file():
            raise RuntimeError("Office tidak menghasilkan PDF. Tutup dialog Office dan periksa dokumen sumber.")
        shutil.move(str(converted), target)
    return {"path": str(target)}


def execute(request: dict, progress=lambda *_: None) -> dict:
    action = request["action"]
    if action == "copy":
        shutil.copyfile(request["source"], request["output"])
        return {"path": request["output"]}
    if action == "render":
        return render(request)
    handlers = {"merge": merge, "split": split, "images_to_pdf": images_to_pdf,
                "pdf_to_images": pdf_to_images, "pdf_to_word": pdf_to_word,
                "pdf_to_excel": pdf_to_excel, "office_to_pdf": office_to_pdf}
    if action in handlers:
        return handlers[action](request, progress)
    return edit(request)
