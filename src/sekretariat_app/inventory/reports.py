from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from importlib.resources import files
from pathlib import Path

from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from sekretariat_app.inventory.models import CardRow, InventoryItem, MutationRow, OutgoingRow


MONTHS_ID = (
    "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
)


def date_id(value: date) -> str:
    return f"{value.day} {MONTHS_ID[value.month]} {value.year}"


def quantity(value: float, *, dash_zero: bool = False) -> str:
    if dash_zero and abs(value) < 1e-9:
        return "-"
    raw = f"{value:,.2f}".rstrip("0").rstrip(".")
    return raw.replace(",", "_").replace(".", ",").replace("_", ".")


def money(value: float, *, dash_zero: bool = False, rupiah: bool = False) -> str:
    if dash_zero and abs(value) < 0.5:
        return "-"
    rendered = f"{round(value):,}".replace(",", ".")
    return f"Rp {rendered}" if rupiah else rendered


class InventoryReportService:
    """Pembuat PDF A4 yang mengikuti tiga contoh fisik pengguna."""

    def __init__(self, settings: dict[str, str]) -> None:
        self.settings = settings
        self.margin = self._number("page_margin_mm", 14.0, 8.0, 25.0) * mm
        self.table_font = self._number("table_font_size", 7.5, 5.5, 10.0)
        self.line_width = self._number("table_line_width", 0.7, 0.3, 1.5)
        base = getSampleStyleSheet()
        self.normal = ParagraphStyle(
            "InventoryNormal", parent=base["Normal"], fontName="Helvetica",
            fontSize=8.5, leading=10.5, textColor=colors.black,
        )
        self.small = ParagraphStyle(
            "InventorySmall", parent=self.normal, fontSize=self.table_font,
            leading=max(self.table_font + 1.2, 7.0),
        )
        self.small_center = ParagraphStyle("InventorySmallCenter", parent=self.small, alignment=TA_CENTER)
        self.small_right = ParagraphStyle("InventorySmallRight", parent=self.small, alignment=TA_RIGHT)
        self.small_bold = ParagraphStyle("InventorySmallBold", parent=self.small, fontName="Helvetica-Bold")
        self.center = ParagraphStyle("InventoryCenter", parent=self.normal, alignment=TA_CENTER)
        self.title = ParagraphStyle(
            "InventoryTitle", parent=self.center, fontName="Helvetica-Bold",
            fontSize=11, leading=13, spaceAfter=2 * mm,
        )

    def _number(self, key: str, default: float, minimum: float, maximum: float) -> float:
        try:
            return max(minimum, min(maximum, float(self.settings.get(key, default))))
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _safe(value: object, *, empty: str = "-") -> str:
        text = str(value or "").strip() or empty
        return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                .replace("\n", "<br/>"))

    def _p(
        self, value: object, style: ParagraphStyle | None = None, *, empty: str = "-",
    ) -> Paragraph:
        return Paragraph(self._safe(value, empty=empty), style or self.small)

    def _document(self, target: str | Path, *, margin: float | None = None) -> SimpleDocTemplate:
        edge = self.margin if margin is None else margin
        return SimpleDocTemplate(
            str(target), pagesize=A4, leftMargin=edge, rightMargin=edge,
            topMargin=10 * mm, bottomMargin=12 * mm,
            title=self.settings.get("office_name", "SEKRETARIAT DPRD"),
            author=self.settings.get("government_name", "PEMERINTAH KOTA BITUNG"),
        )

    def _logo(self) -> Image | Spacer:
        configured = Path(self.settings.get("logo_path", "")).expanduser()
        source: Path | None = configured if configured.is_file() else None
        if source is None:
            try:
                bundled = files("sekretariat_app.resources").joinpath("app_icon.ico")
                source = Path(str(bundled))
            except (ModuleNotFoundError, TypeError):
                source = None
        if source is None or not source.is_file():
            return Spacer(27 * mm, 24 * mm)
        try:
            with PILImage.open(source) as original:
                image = original.convert("RGBA")
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                buffer.seek(0)
            return Image(buffer, width=27 * mm, height=24 * mm, kind="proportional")
        except OSError:
            return Spacer(27 * mm, 24 * mm)

    def _letterhead(self) -> list[object]:
        office = self._safe(self.settings.get("office_name"), empty="")
        government = self._safe(self.settings.get("government_name"), empty="")
        address = self._safe(self.settings.get("office_address"), empty="")
        website = self._safe(self.settings.get("office_website"), empty="")
        postal = self._safe(self.settings.get("office_postal_code"), empty="")
        heading = Paragraph(
            f"<b><font size='15'>{government}</font></b><br/>"
            f"<b><font size='18'>{office}</font></b><br/>"
            f"<font size='9'>{address}</font><br/>"
            f"<font size='9'><u>{website}</u> &nbsp; {postal}</font>",
            ParagraphStyle("Letterhead", parent=self.center, leading=17),
        )
        table = Table([[self._logo(), heading, Spacer(5 * mm, 1)]], colWidths=[30 * mm, 130 * mm, 5 * mm])
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1 * mm),
        ]))
        return [
            table,
            HRFlowable(width="100%", thickness=2.0, color=colors.black, spaceBefore=0.5 * mm),
            HRFlowable(width="100%", thickness=0.6, color=colors.black, spaceBefore=0.4 * mm),
            Spacer(1, 3.5 * mm),
        ]

    def _signature_table(self, document_date: date, *, show_city: bool = False) -> Table:
        left_title = self._safe(self.settings.get("secretary_title"))
        left_name = self._safe(self.settings.get("secretary_name"))
        left_nip = self._safe(self.settings.get("secretary_nip"))
        right_title = self._safe(self.settings.get("manager_title"))
        right_name = self._safe(self.settings.get("manager_name"))
        right_nip = self._safe(self.settings.get("manager_nip"))
        city = self._safe(self.settings.get("city", "Bitung"))
        left = Paragraph(
            f"<b>MENGETAHUI<br/>{left_title}</b><br/><br/><br/><br/>"
            f"<b><u>{left_name}</u></b><br/>NIP. {left_nip}", self.center,
        )
        right_prefix = f"{city}, {date_id(document_date)}<br/><br/>" if show_city else ""
        right = Paragraph(
            f"{right_prefix}<b>{right_title}</b><br/><br/><br/><br/>"
            f"<b><u>{right_name}</u></b><br/>NIP. {right_nip}", self.center,
        )
        table = Table([[left, right]], colWidths=[84 * mm, 84 * mm])
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ]))
        return table

    def build_issue_order(self, target: str | Path, document: dict[str, object]) -> Path:
        header = document["header"]
        lines = list(document["lines"])
        document_date = date.fromisoformat(str(header["book_date"]))
        story: list[object] = self._letterhead()
        story.extend((
            Paragraph(self._safe(self.settings.get("issue_title")), self.title),
            Paragraph(f"<b>NO. {self._safe(header['document_number'])}</b>", self.center),
            Spacer(1, 5 * mm),
        ))
        info = Table([
            [self._p("Dari", self.normal), self._p(":", self.normal), self._p(header["sender"], self.normal)],
            [self._p("Kepada", self.normal), self._p(":", self.normal), self._p(header["recipient"], self.normal)],
            [self._p("Alamat", self.normal), self._p(":", self.normal), self._p(header["address"], self.normal)],
        ], colWidths=[16 * mm, 4 * mm, 145 * mm])
        info.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        story.append(info)
        division = self._safe(header["division"])
        office_full_name = self._safe(self.settings.get("office_full_name", "Sekretariat DPRD Kota Bitung"))
        story.append(Paragraph(
            "Harap dikeluarkan dari gudang dan disalurkan barang tersebut dalam daftar "
            f"di bawah ini untuk keperluan {division} {office_full_name}.", self.normal,
        ))
        if header.get("request_basis"):
            story.append(Paragraph(self._safe(header["request_basis"]), self.normal))
        story.append(Spacer(1, 4 * mm))

        rows: list[list[object]] = [[
            self._p("No Urt.", self.small_center), self._p("Banyaknya", self.small_center),
            self._p("Nama Barang", self.small_center), self._p("Harga Satuan", self.small_center),
            self._p("Jumlah", self.small_center), self._p("Ket.", self.small_center),
        ]]
        for index, line in enumerate(lines, start=1):
            rows.append([
                self._p(index, self.small_center),
                self._p(f"{quantity(float(line['quantity']))} {line['unit']}", self.small_center),
                self._p(line["name"]),
                self._p(money(float(line["unit_price"])), self.small_right),
                self._p(money(float(line["total_cost"])), self.small_right),
                self._p(line.get("note", ""), self.small_center, empty=""),
            ])
        try:
            blank_rows = max(0, min(25, int(self.settings.get("issue_blank_rows", "8"))))
        except ValueError:
            blank_rows = 8
        for _ in range(blank_rows):
            rows.append(["", "", "", "", "", ""])
        table = Table(rows, colWidths=[14 * mm, 27 * mm, 57 * mm, 25 * mm, 25 * mm, 22 * mm], repeatRows=1)
        table.setStyle(self._grid_style(header_rows=1))
        story.extend((table, Spacer(1, 5 * mm), self._signature_table(document_date, show_city=True)))
        self._document(target).build(story)
        return Path(target)

    def build_card(
        self, target: str | Path, item: InventoryItem, rows: list[CardRow],
    ) -> Path:
        doc = self._document(target, margin=11 * mm)
        metadata = Table([
            [self._p("SKPD", self.normal), self._p(":", self.normal), self._p(self.settings.get("office_full_name", "Sekretariat DPRD Kota Bitung").upper(), self.normal), "", self._p(self.settings.get("card_attachment"), self.normal)],
            [self._p("KAB/KOTA", self.normal), self._p(":", self.normal), self._p("KOTA BITUNG", self.normal), "", ""],
            [self._p("PROVINSI", self.normal), self._p(":", self.normal), self._p(self.settings.get("province"), self.normal), "", ""],
        ], colWidths=[18 * mm, 4 * mm, 78 * mm, 44 * mm, 33 * mm])
        metadata.setStyle(TableStyle([
            ("SPAN", (4, 0), (4, 2)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (4, 0), (4, 2), "CENTER"), ("FONTNAME", (4, 0), (4, 0), "Helvetica-Bold"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        story: list[object] = [metadata, Spacer(1, 1 * mm), Paragraph(self._safe(self.settings.get("card_title")), self.title)]
        detail = Table([
            [self._p("Gudang", self.normal), self._p(":", self.normal), self._p(self.settings.get("warehouse"), self.normal, empty=""), "", self._p("Kartu No", self.normal), self._p(":", self.normal), self._p(item.card_number, self.normal, empty="")],
            [self._p("Nama Barang", self.normal), self._p(":", self.normal), self._p(item.name, self.normal), self._p(f"Rp. {money(item.standard_price)}", self.normal), self._p("Spesifikasi", self.normal), self._p(":", self.normal), self._p(item.specification, self.normal, empty="")],
            [self._p("Satuan", self.normal), self._p(":", self.normal), self._p(item.unit, self.normal), "", "", "", ""],
        ], colWidths=[23 * mm, 4 * mm, 65 * mm, 30 * mm, 22 * mm, 4 * mm, 29 * mm])
        detail.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ("FONTNAME", (2, 1), (2, 2), "Helvetica-Bold"),
        ]))
        story.extend((detail, Spacer(1, 4 * mm)))
        table_rows: list[list[object]] = [
            [self._p("Tanggal", self.small_center), self._p("No Surat<br/>Penerimaan/Pengeluaran", self.small_center), self._p("Uraian", self.small_center), self._p("Barang-barang", self.small_center), "", "", self._p("Harga<br/>Satuan", self.small_center), self._p("Jumlah Harga Barang yang Diterima/<br/>Yang Dikeluarkan/Sisa", self.small_center), "", "", self._p("Ket", self.small_center)],
            ["", "", "", self._p("Masuk", self.small_center), self._p("Keluar", self.small_center), self._p("Sisa", self.small_center), "", self._p("Bertambah", self.small_center), self._p("Berkurang", self.small_center), self._p("Sisa", self.small_center), ""],
            [
                self._p("1", self.small_center), self._p("2", self.small_center), "",
                self._p("3", self.small_center), self._p("4", self.small_center),
                self._p("5", self.small_center), self._p("6", self.small_center),
                self._p("7", self.small_center), self._p("8", self.small_center),
                self._p("9", self.small_center), self._p("10", self.small_center),
            ],
        ]
        for row in rows:
            table_rows.append([
                self._p(row.book_date.strftime("%d-%m-%Y"), self.small_center),
                self._p(row.document_number, self.small_center), self._p(row.description, self.small_center),
                self._p(quantity(row.quantity_in, dash_zero=True), self.small_center),
                self._p(quantity(row.quantity_out, dash_zero=True), self.small_center),
                self._p(quantity(row.balance_quantity, dash_zero=True), self.small_center),
                self._p(money(row.unit_price, dash_zero=True), self.small_right),
                self._p(money(row.value_in, dash_zero=True), self.small_right),
                self._p(money(row.value_out, dash_zero=True), self.small_right),
                self._p(money(row.balance_value, dash_zero=True), self.small_right),
                self._p(row.note, self.small_center, empty=""),
            ])
        table = Table(
            table_rows,
            colWidths=[17 * mm, 31 * mm, 23 * mm, 9 * mm, 9 * mm, 9 * mm,
                       16 * mm, 19 * mm, 19 * mm, 19 * mm, 9 * mm],
            repeatRows=3,
        )
        style = self._grid_style(header_rows=3)
        style.add("SPAN", (0, 0), (0, 1)); style.add("SPAN", (1, 0), (1, 1))
        style.add("SPAN", (2, 0), (2, 1)); style.add("SPAN", (3, 0), (5, 0))
        style.add("SPAN", (6, 0), (6, 1)); style.add("SPAN", (7, 0), (9, 0))
        style.add("SPAN", (10, 0), (10, 1))
        table.setStyle(style)
        story.append(table)
        doc.build(story, onFirstPage=self._page_number, onLaterPages=self._page_number)
        return Path(target)

    def build_mutation(self, target: str | Path, rows: list[MutationRow], as_of: date) -> Path:
        story: list[object] = self._letterhead()
        story.extend((
            Spacer(1, 2 * mm), Paragraph(self._safe(self.settings.get("mutation_title")), self.title),
            Paragraph(f"S/D {date_id(as_of).upper()}", self.center), Spacer(1, 5 * mm),
        ))
        previous_year = as_of.year - 1
        table_rows: list[list[object]] = [
            [self._p("NO", self.small_center), self._p("URAIAN", self.small_center), self._p(f"NILAI S/D DESEMBER {previous_year}", self.small_center), "", "", self._p("MUTASI", self.small_center), "", self._p(f"NILAI S/D {MONTHS_ID[as_of.month].upper()} {as_of.year}", self.small_center), "", ""],
            ["", "", self._p("JUMLAH", self.small_center), self._p("SATUAN", self.small_center), self._p("RUPIAH", self.small_center), self._p("TAMBAH", self.small_center), self._p("KURANG", self.small_center), self._p("JUMLAH", self.small_center), self._p("SATUAN", self.small_center), self._p("RUPIAH", self.small_center)],
        ]
        spans: list[tuple] = [
            ("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (1, 1)),
            ("SPAN", (2, 0), (4, 0)), ("SPAN", (5, 0), (6, 0)),
            ("SPAN", (7, 0), (9, 0)),
        ]
        categories: dict[str, list[MutationRow]] = defaultdict(list)
        for row in rows:
            categories[row.category or "Lainnya"].append(row)
        sequence = 0
        for category_index, (category, members) in enumerate(categories.items()):
            category_row = len(table_rows)
            letter = chr(ord("A") + category_index) if category_index < 26 else str(category_index + 1)
            table_rows.append([self._p(letter, self.small_center), self._p(category, self.small_bold), "", "", "", "", "", "", "", ""])
            spans.append(("SPAN", (1, category_row), (9, category_row)))
            for row in members:
                sequence += 1
                table_rows.append([
                    self._p(sequence, self.small_center), self._p(row.name),
                    self._p(quantity(row.opening_quantity, dash_zero=True), self.small_right),
                    self._p(money(row.opening_price, dash_zero=True), self.small_right),
                    self._p(money(row.opening_value, dash_zero=True), self.small_right),
                    self._p(quantity(row.quantity_in, dash_zero=True), self.small_right),
                    self._p(quantity(row.quantity_out, dash_zero=True), self.small_right),
                    self._p(quantity(row.closing_quantity, dash_zero=True), self.small_right),
                    self._p(money(row.closing_price, dash_zero=True), self.small_right),
                    self._p(money(row.closing_value, dash_zero=True), self.small_right),
                ])
        table = Table(
            table_rows,
            colWidths=[8 * mm, 41 * mm, 14 * mm, 16 * mm, 19 * mm,
                       13 * mm, 13 * mm, 14 * mm, 17 * mm, 22 * mm],
            repeatRows=2,
        )
        style = self._grid_style(header_rows=2)
        for command in spans:
            style.add(*command)
        table.setStyle(style)
        story.extend((table, Spacer(1, 7 * mm), KeepTogether([self._signature_table(as_of)])))
        self._document(target).build(story)
        return Path(target)

    def build_outgoing_recap(
        self, target: str | Path, rows: list[OutgoingRow], start: date,
        end: date, division: str = "",
    ) -> Path:
        story: list[object] = self._letterhead()
        subtitle = f"Periode {date_id(start)} s/d {date_id(end)}"
        if division:
            subtitle += f"<br/>{self._safe(division)}"
        story.extend((
            Paragraph("LAPORAN REKAPAN BARANG KELUAR", self.title),
            Paragraph(subtitle, self.center), Spacer(1, 5 * mm),
        ))
        table_rows: list[list[object]] = [[
            self._p("No", self.small_center), self._p("Tanggal", self.small_center),
            self._p("No. Surat", self.small_center), self._p("Nama Barang", self.small_center),
            self._p("Jumlah", self.small_center), self._p("Harga Satuan", self.small_center),
            self._p("Nilai", self.small_center), self._p("Bagian", self.small_center),
            self._p("Keterangan", self.small_center),
        ]]
        for index, row in enumerate(rows, start=1):
            table_rows.append([
                self._p(index, self.small_center), self._p(row.book_date.strftime("%d-%m-%Y"), self.small_center),
                self._p(row.document_number, self.small_center), self._p(row.name),
                self._p(f"{quantity(row.quantity)} {row.unit}", self.small_center),
                self._p(money(row.unit_price), self.small_right), self._p(money(row.total_cost), self.small_right),
                self._p(row.division), self._p(row.note or row.description),
            ])
        table = Table(table_rows, colWidths=[7 * mm, 17 * mm, 24 * mm, 32 * mm, 17 * mm, 20 * mm, 21 * mm, 31 * mm, 21 * mm], repeatRows=1)
        table.setStyle(self._grid_style(header_rows=1))
        total = sum(row.total_cost for row in rows)
        story.extend((table, Spacer(1, 3 * mm), Paragraph(f"<b>Total Nilai Barang Keluar: Rp {money(total)}</b>", self.small_right)))
        self._document(target, margin=9 * mm).build(story)
        return Path(target)

    def _grid_style(self, *, header_rows: int) -> TableStyle:
        return TableStyle([
            ("GRID", (0, 0), (-1, -1), self.line_width, colors.black),
            ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
            ("BACKGROUND", (0, 0), (-1, header_rows - 1), colors.white),
            ("ALIGN", (0, 0), (-1, header_rows - 1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 1.2 * mm),
            ("RIGHTPADDING", (0, 0), (-1, -1), 1.2 * mm),
            ("TOPPADDING", (0, 0), (-1, -1), 0.8 * mm),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0.8 * mm),
            ("LEADING", (0, 0), (-1, -1), max(self.table_font + 1, 7)),
        ])

    @staticmethod
    def _page_number(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawCentredString(A4[0] / 2, 6 * mm, f"Page {document.page}")
        canvas.restoreState()


def export_outgoing_xlsx(
    target: str | Path, rows: list[OutgoingRow], start: date, end: date, division: str,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rekap Barang Keluar"
    sheet.append(["LAPORAN REKAPAN BARANG KELUAR"])
    sheet.append([f"Periode {date_id(start)} s/d {date_id(end)}"])
    sheet.append([division or "Semua Bagian"])
    sheet.append([])
    headers = ["No", "Tanggal", "No. Surat", "Kode", "Nama Barang", "Jumlah", "Satuan", "Harga Satuan", "Nilai", "Bagian", "Keterangan"]
    sheet.append(headers)
    for index, row in enumerate(rows, start=1):
        sheet.append([
            index, row.book_date, row.document_number, row.code, row.name,
            row.quantity, row.unit, row.unit_price, row.total_cost, row.division,
            row.note or row.description,
        ])
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 8, "TOTAL")
    sheet.cell(total_row, 9, f"=SUM(I6:I{total_row - 1})")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    for row_index in (1, 2, 3):
        sheet.cell(row_index, 1).alignment = Alignment(horizontal="center")
        sheet.cell(row_index, 1).font = Font(bold=True, size=14 if row_index == 1 else 10)
    thin = Side(style="thin", color="000000")
    for cell in sheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=5, max_row=total_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=6, max_row=total_row - 1, min_col=8, max_col=9):
        for cell in row:
            cell.number_format = '#,##0'
    sheet.cell(total_row, 9).number_format = '#,##0'
    widths = [6, 13, 24, 14, 32, 12, 12, 16, 17, 40, 34]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:K{total_row - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(target)
    return Path(target)
